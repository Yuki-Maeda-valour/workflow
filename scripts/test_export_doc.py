"""export-doc の出力前検査と CSV/xlsx の正常対照。"""
import contextlib
import importlib.util
import io
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

SCRIPT = Path(__file__).resolve().parents[1] / 'plugins/dev-workflow/skills/export-doc/scripts/md_tables_to_xlsx.py'
HAS_OPENPYXL = importlib.util.find_spec('openpyxl') is not None
TABLE = '# Heading\n\n| Name | Value |\n| --- | --- |\n| 日本語 | **bold** |\n'


def module():
    spec = importlib.util.spec_from_file_location('export_tables', SCRIPT)
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    return helper


class ExportDocTest(unittest.TestCase):
    def setUp(self):
        self.helper = module()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.doc = self.root / 'doc'
        self.doc.mkdir()
        self.source = self.doc / 'input.md'
        self.source.write_text(TABLE)
        self.out = self.root / 'export' / 'book.xlsx'

    def test_csv_final_symlink_never_changes_source(self):
        self.out.with_suffix('').mkdir(parents=True)
        target = self.out.with_suffix('') / 'Sheet.csv'
        target.symlink_to(self.source)
        before = self.source.read_bytes()
        with contextlib.suppress(ValueError):
            self.helper.write_csv_fallback([('Sheet', [['overwrite']])], self.out)
        self.assertEqual(before, self.source.read_bytes())
        self.assertTrue(target.is_symlink())

    def cli(self, *args, csv=False):
        output = io.StringIO()
        with contextlib.ExitStack() as stack:
            stack.enter_context(mock.patch.object(sys, 'argv', [str(SCRIPT), *map(str, args)]))
            stack.enter_context(contextlib.redirect_stdout(output))
            stack.enter_context(contextlib.redirect_stderr(output))
            if csv:
                stack.enter_context(mock.patch.dict(sys.modules, {'openpyxl': None}))
            rc = self.helper.main()
        return rc, output.getvalue()

    def test_normal_cli_initial_and_overwrite_both_formats(self):
        for csv_mode in (True, False):
            with self.subTest(csv=csv_mode):
                if not csv_mode and not HAS_OPENPYXL:
                    self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                out = self.root / ('space csv' if csv_mode else 'space xlsx') / 'book.xlsx'
                for _ in range(2):
                    rc, message = self.cli(self.source, '-o', out, csv=csv_mode)
                    self.assertEqual(0, rc, message)
                if csv_mode:
                    data = (out.with_suffix('') / 'input Heading.csv').read_bytes()
                    self.assertTrue(data.startswith(b'\xef\xbb\xbf'))
                    self.assertIn('日本語,bold', data.decode('utf-8-sig'))
                    self.assertIn('FALLBACK', message)
                else:
                    import openpyxl
                    wb = openpyxl.load_workbook(out)
                    self.addCleanup(wb.close)
                    self.assertEqual(['input Heading'], wb.sheetnames)
                    ws = wb.active
                    self.assertEqual('日本語', ws['A2'].value)
                    self.assertEqual('bold', ws['B2'].value)
                    self.assertTrue(ws['A1'].font.bold)
                    self.assertEqual('A2', ws.freeze_panes)

    def test_source_directory_rejected_before_mkdir(self):
        for csv_mode in (True, False):
            with self.subTest(csv=csv_mode):
                if not csv_mode and not HAS_OPENPYXL:
                    self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                out = self.doc / 'new' / 'book.xlsx'
                rc, message = self.cli(self.source, '-o', out, csv=csv_mode)
                self.assertEqual(2, rc, message)
                self.assertIn('ERROR', message)
                self.assertFalse((self.doc / 'new').exists())
                self.assertEqual(TABLE, self.source.read_text())

    @unittest.skipUnless(HAS_OPENPYXL, 'openpyxl 未導入: xlsx の入力直接重複は未確認')
    def test_direct_input_rejected_even_with_explicit_other_root(self):
        other = self.root / 'other'
        other.mkdir()
        rc, message = self.cli(self.source, '-o', self.source, '--source-root', other)
        self.assertEqual(2, rc, message)
        self.assertEqual(TABLE, self.source.read_text())

    def test_scratch_and_multiple_source_roots(self):
        second = self.root / 'second doc'
        second.mkdir()
        extra = second / 'other.md'
        extra.write_text(TABLE)
        scratch = self.root / 'scratch'
        scratch.mkdir()
        copied = scratch / 'copy.md'
        copied.write_text(TABLE)
        for root in (self.doc, second):
            for csv_mode in (True, False):
                with self.subTest(csv=csv_mode):
                    if not csv_mode and not HAS_OPENPYXL:
                        self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                    rc, message = self.cli(copied, extra, '-o', root / 'new' / 'book.xlsx',
                                           '--source-root', self.doc, '--source-root', second, csv=csv_mode)
                    self.assertEqual(2, rc, message)
                    self.assertFalse((root / 'new').exists())
        rc, message = self.cli(copied, extra, '-o', self.out, '--source-root', self.doc,
                               '--source-root', second)
        self.assertEqual(0, rc, message)

    def test_parent_symlink_safe_and_source_overlap(self):
        for name, target, expected in [('safe', self.root / 'space output', 0), ('bad', self.doc, 2)]:
            target.mkdir(exist_ok=True)
            alias = self.root / name
            alias.symlink_to(target, target_is_directory=True)
            for csv_mode in (True, False):
                with self.subTest(csv=csv_mode):
                    if not csv_mode and not HAS_OPENPYXL:
                        self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                    rc, message = self.cli(self.source, '-o', alias / 'new' / 'book.xlsx', csv=csv_mode)
                    self.assertEqual(expected, rc, message)
            if expected:
                self.assertFalse((target / 'new').exists())

    def test_final_symlink_and_hardlink_rejected_in_both_formats(self):
        # Both the copied source and an unselected document must survive.
        scratch = self.root / 'scratch'
        scratch.mkdir()
        copied = scratch / 'copy.md'
        copied.write_text(TABLE)
        unselected = self.doc / 'unselected.md'
        unselected.write_text('private original')
        for original in (self.source, unselected):
            for kind in ('symlink', 'hardlink', 'dangling'):
                for csv_mode in (True, False):
                    with self.subTest(original=original.name, kind=kind, csv=csv_mode):
                        if not csv_mode and not HAS_OPENPYXL:
                            self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                        out = self.root / f'{original.stem}-{kind}-{csv_mode}' / 'book.xlsx'
                        dest = out.with_suffix('') / 'copy Heading.csv' if csv_mode else out
                        dest.parent.mkdir(parents=True)
                        if kind == 'hardlink':
                            os.link(original, dest)
                        else:
                            dest.symlink_to(original if kind == 'symlink' else self.root / 'missing')
                        before = original.read_bytes()
                        rc, message = self.cli(copied, '-o', out, '--source-root', self.doc, csv=csv_mode)
                        self.assertEqual(2, rc, message)
                        self.assertEqual(before, original.read_bytes())
                        self.assertTrue(dest.is_symlink() if kind != 'hardlink' else dest.exists())
                        dest.unlink()

    def test_all_csv_outputs_preflight_before_overwrite(self):
        for kind in ('symlink', 'hardlink', 'fifo'):
            with self.subTest(kind=kind):
                out = self.root / kind / 'book.xlsx'
                directory = out.with_suffix('')
                directory.mkdir(parents=True)
                first = directory / 'first.csv'
                first.write_text('keep first')
                last = directory / 'last.csv'
                if kind == 'symlink':
                    last.symlink_to(self.source)
                elif kind == 'hardlink':
                    os.link(self.source, last)
                else:
                    os.mkfifo(last)
                with self.assertRaises(self.helper.OutputError):
                    self.helper.write_csv_fallback([('first', [['one']]), ('last', [['two']])], out)
                self.assertEqual('keep first', first.read_text())
                self.assertEqual(TABLE, self.source.read_text())
                last.unlink()

    def test_broken_cyclic_and_file_parents(self):
        broken = self.root / 'broken'
        broken.symlink_to(self.root / 'missing')
        cycle = self.root / 'cycle'
        cycle.symlink_to(cycle)
        regular = self.root / 'regular'
        regular.write_text('keep parent')
        for parent in (broken, cycle, regular):
            for csv_mode in (True, False):
                with self.subTest(csv=csv_mode):
                    if not csv_mode and not HAS_OPENPYXL:
                        self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                    rc, message = self.cli(self.source, '-o', parent / 'child' / 'book.xlsx', csv=csv_mode)
                    self.assertEqual(2, rc, message)
                    self.assertIn('ERROR', message)
        self.assertFalse((self.root / 'missing').exists())
        self.assertEqual('keep parent', regular.read_text())

    def test_non_regular_xlsx_and_direct_xlsx_symlink(self):
        for kind in ('directory', 'fifo', 'symlink'):
            out = self.root / kind
            if kind == 'directory':
                out.mkdir()
            elif kind == 'fifo':
                os.mkfifo(out)
            else:
                out.symlink_to(self.source)
            with self.assertRaises(self.helper.OutputError):
                self.helper.write_xlsx([('Sheet', [['value']])], out)
        self.assertEqual(TABLE, self.source.read_text())

    def test_invalid_source_root_rejected(self):
        for root in (self.root / 'missing', self.source):
            rc, message = self.cli(self.source, '-o', self.out, '--source-root', root)
            self.assertEqual(2, rc, message)
            self.assertFalse(self.out.parent.exists())

    def test_symlink_input_protects_both_selected_and_real_parents(self):
        selected = self.root / 'selected'
        selected.mkdir()
        link = selected / 'alias.md'
        link.symlink_to(self.source)
        for root in (selected, self.doc):
            for csv_mode in (True, False):
                with self.subTest(csv=csv_mode):
                    if not csv_mode and not HAS_OPENPYXL:
                        self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                    rc, message = self.cli(link, '-o', root / 'new' / 'book.xlsx', csv=csv_mode)
                    self.assertEqual(2, rc, message)
                    self.assertFalse((root / 'new').exists())

    def test_missing_and_no_tables_behavior(self):
        empty = self.doc / 'empty.md'
        empty.write_text('# no table')
        for csv_mode in (True, False):
            with self.subTest(csv=csv_mode):
                if not csv_mode and not HAS_OPENPYXL:
                    self.skipTest('openpyxl 未導入: xlsx 経路は未確認')
                self.out = self.root / str(csv_mode) / 'book.xlsx'
                self.assertEqual(2, self.cli(self.root / 'missing', '-o', self.out, csv=csv_mode)[0])
                self.assertEqual(1, self.cli(empty, '-o', self.out, csv=csv_mode)[0])
                self.assertFalse(self.out.parent.exists())
                self.assertEqual(0, self.cli(self.root / 'missing', self.source, '-o', self.out, csv=csv_mode)[0])


if __name__ == '__main__':
    unittest.main()
